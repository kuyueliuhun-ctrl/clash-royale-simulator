#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Markdown → DOCX 转换器（离线、无外部依赖，仅用 python-docx）。

支持：`#`~`######` 标题（映射 Word Heading 1~6）、段落、`-`/`*` 无序列表、
`1.` 有序列表、``` 围栏代码块、`| a | b |` 表格（首个分隔行必须形如 `|---|`）、
`>` 引用、`---` 分隔线、行内 `**粗体**` / `` `代码` `` / `[文本](链接)`。
另在文首插入**静态目录**（由标题自动生成）与一个 Word TOC 域（打开后可按 F9 更新）。

用法：
    PYTHONIOENCODING=utf-8 .venv/Scripts/python.exe scripts/_survey_md_to_docx.py \
        --md docs/training_method.md --docx docs/training_method.docx \
        --title "训练方法文档"
"""
from __future__ import annotations

import argparse
import io
import re
import sys
from typing import List, Tuple

# T1-1b 补齐（2026-09-19）：原先这里是**手写**的 UTF-8 兜底块（只处理 stdout、且不处理 stderr
# ⇒ traceback 在 GBK 下仍是乱码）。现收敛到 T1-1 的**单一实现**（两路 + errors='replace'）。
import os as _os
import sys as _sys
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))),
                                 "src", "clasher_new"))
from rl.io_bootstrap import force_utf8_stdout  # noqa: E402
force_utf8_stdout()

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

BODY_FONT = "微软雅黑"
CODE_FONT = "Consolas"
BODY_SIZE = Pt(10.5)
CODE_SIZE = Pt(9)
TABLE_SIZE = Pt(9)

H_RE = re.compile(r"^(#{1,6})\s+(.*)$")
UL_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
OL_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")


# ---------------------------------------------------------------- low level
def set_run_font(run, name: str, size: Pt, bold=None, italic=None, color=None):
    run.font.name = name
    run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    rf.set(qn("w:eastAsia"), name)


def shade(element, hex_fill: str):
    pr = element.get_or_add_pPr() if element.tag.endswith("}p") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def cell_shade(cell, hex_fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_toc_field(paragraph):
    """插入 Word TOC 域（打开文档后 F9 更新即可显示页码）。"""
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "（Word 目录域：在 Word 中按 F9 或右键“更新域”生成带页码目录）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (fld, instr, sep, t, end):
        run._element.append(el)


# ---------------------------------------------------------------- inline
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, base_size=BODY_SIZE, base_bold=False):
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            r = paragraph.add_run(chunk[2:-2])
            set_run_font(r, BODY_FONT, base_size, bold=True)
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            r = paragraph.add_run(chunk[1:-1])
            set_run_font(r, CODE_FONT, Pt(base_size.pt - 0.5), color=RGBColor(0xB0, 0x30, 0x00))
        elif chunk.startswith("[") and "](" in chunk:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", chunk)
            r = paragraph.add_run(f"{m.group(1)}（{m.group(2)}）")
            set_run_font(r, BODY_FONT, base_size, color=RGBColor(0x0B, 0x53, 0x94))
        else:
            r = paragraph.add_run(chunk)
            set_run_font(r, BODY_FONT, base_size, bold=base_bold)


# ---------------------------------------------------------------- conversion
def configure_styles(doc: Document):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = BODY_SIZE
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15
    for i in range(1, 7):
        h = doc.styles[f"Heading {i}"]
        h.font.name = BODY_FONT
        h.font.size = Pt(max(11, 20 - i * 1.6))
        h.font.bold = True
        h.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)


def split_table_row(line: str) -> List[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", s)]


def convert(md_path: str, docx_path: str, title: str | None = None) -> dict:
    with io.open(md_path, encoding="utf-8", errors="replace") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin = sec.bottom_margin = Cm(1.9)

    # 标题页
    t = title or (lines[0][2:].strip() if lines and lines[0].startswith("# ") else "文档")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    set_run_font(r, BODY_FONT, Pt(26), bold=True, color=RGBColor(0x1F, 0x38, 0x64))
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"源文件：{md_path}")
    set_run_font(r2, CODE_FONT, Pt(9), color=RGBColor(0x60, 0x60, 0x60))
    doc.add_paragraph()

    # 静态目录（Heading 1~3）
    hp = doc.add_paragraph()
    r = hp.add_run("目录（静态，章节号与正文一致）")
    set_run_font(r, BODY_FONT, Pt(15), bold=True)
    headings: List[Tuple[int, str]] = []
    for ln in lines:
        m = H_RE.match(ln)
        if m and len(m.group(1)) <= 3:
            headings.append((len(m.group(1)), m.group(2).strip()))
    for lvl, txt in headings:
        tp = doc.add_paragraph()
        tp.paragraph_format.left_indent = Cm(0.5 * (lvl - 1))
        tp.paragraph_format.space_after = Pt(1)
        rr = tp.add_run(txt)
        set_run_font(rr, BODY_FONT, Pt(11 - lvl * 0.6), bold=(lvl == 1))
    doc.add_paragraph()
    tp = doc.add_paragraph()
    rr = tp.add_run("（以下为 Word 自动目录域，可在 Word 中按 F9 更新出页码）")
    set_run_font(rr, BODY_FONT, Pt(9), italic=True, color=RGBColor(0x60, 0x60, 0x60))
    add_toc_field(doc.add_paragraph())
    doc.add_page_break()

    i = 0
    n_tables = n_code = n_par = 0
    first_h1_skipped = False
    while i < len(lines):
        ln = lines[i]
        stripped = ln.strip()

        # 代码块
        if stripped.startswith("```"):
            i += 1
            buf: List[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            cp = doc.add_paragraph()
            cp.paragraph_format.space_before = Pt(3)
            cp.paragraph_format.space_after = Pt(6)
            cp.paragraph_format.left_indent = Cm(0.35)
            shade(cp._element, "F3F4F6")
            r = cp.add_run("\n".join(buf) if buf else " ")
            set_run_font(r, CODE_FONT, CODE_SIZE)
            n_code += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1].strip()):
            header = split_table_row(lines[i])
            i += 2
            rows: List[List[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            ncol = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
            tb = doc.add_table(rows=1, cols=ncol)
            tb.style = "Table Grid"
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c in range(ncol):
                cell = tb.rows[0].cells[c]
                cell.text = ""
                para = cell.paragraphs[0]
                add_inline(para, header[c] if c < len(header) else "", TABLE_SIZE, True)
                cell_shade(cell, "DCE6F1")
            for row in rows:
                cells = tb.add_row().cells
                for c in range(ncol):
                    cells[c].text = ""
                    add_inline(cells[c].paragraphs[0], row[c] if c < len(row) else "", TABLE_SIZE)
            doc.add_paragraph()
            n_tables += 1
            continue

        m = H_RE.match(ln)
        if m:
            lvl = len(m.group(1))
            txt = m.group(2).strip()
            if lvl == 1 and not first_h1_skipped and txt == t:
                first_h1_skipped = True
                i += 1
                continue
            h = doc.add_heading(level=lvl)
            add_inline(h, txt, Pt(max(11, 20 - lvl * 1.6)), True)
            for r in h.runs:
                set_run_font(
                    r,
                    BODY_FONT,
                    Pt(max(11, 20 - lvl * 1.6)),
                    bold=True,
                    color=RGBColor(0x1F, 0x38, 0x64),
                )
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, stripped[2:], Pt(10))
            for r in p.runs:
                r.font.italic = True
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            p = doc.add_paragraph()
            pPr = p._element.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bt = OxmlElement("w:bottom")
            bt.set(qn("w:val"), "single")
            bt.set(qn("w:sz"), "6")
            bt.set(qn("w:color"), "AAAAAA")
            bd.append(bt)
            pPr.append(bd)
            i += 1
            continue

        mo = OL_RE.match(ln)
        if mo:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, mo.group(3))
            n_par += 1
            i += 1
            continue

        mu = UL_RE.match(ln)
        if mu:
            p = doc.add_paragraph(style="List Bullet")
            indent = len(mu.group(1)) // 2
            p.paragraph_format.left_indent = Cm(0.75 + 0.5 * indent)
            add_inline(p, mu.group(2))
            n_par += 1
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, ln.strip())
        n_par += 1
        i += 1

    doc.save(docx_path)
    return {
        "md": md_path,
        "docx": docx_path,
        "headings": len(headings),
        "paragraphs": n_par,
        "tables": n_tables,
        "code_blocks": n_code,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--docx", required=True)
    ap.add_argument("--title", default=None)
    args = ap.parse_args()
    info = convert(args.md, args.docx, args.title)
    print("[md->docx] " + ", ".join(f"{k}={v}" for k, v in info.items()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
